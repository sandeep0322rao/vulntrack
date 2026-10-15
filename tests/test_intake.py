from datetime import timedelta
from io import BytesIO
import re
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from vulntrack.extensions import db
from vulntrack.models import ActiveScanApproval, Finding, ImportBatch, ScanJob, now
from vulntrack.services.spreadsheet_import import spreadsheet_records


def workbook_bytes(rows, sheet_name='Findings'):
    workbook = Workbook()
    workbook.active.title = sheet_name
    for row in rows:
        workbook.active.append(row)
    result = BytesIO()
    workbook.save(result)
    return result.getvalue()


def upload(client, assessment, payload, name='findings.xlsx'):
    return client.post(f'/assessments/{assessment.id}/upload-report', data={'report': (BytesIO(payload), name), 'scope_confirmed': 'y'})


def test_guided_three_paths_and_template(signed_in, assessment):
    response = signed_in.get(f'/assessments/{assessment.id}/add-findings')
    assert response.status_code == 200
    for label in [b'Scan a website', b'Upload a report', b'Add manually']:
        assert label in response.data
    assert response.data.count(b'class="intake-choice"') == 3
    for path in [f'/assessments/{assessment.id}/automatic-scan', f'/assessments/{assessment.id}/upload-report', f'/assessments/{assessment.id}/findings/new']:
        assert signed_in.get(path).status_code == 200
    response = signed_in.get('/templates/findings.xlsx')
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    assert workbook['Findings']['A1'].value == 'Title'
    assert workbook['Findings']['B1'].value == 'Severity'
    assert workbook['Findings'].data_validations.count == 1


def test_excel_preview_then_atomic_idempotent_confirm(signed_in, assessment):
    payload = workbook_bytes([['Title','Severity','URL','Description','Status'], ['Missing CSP','low',assessment.target_url,'Header observation','remediated'], ['Missing CSP','Low',assessment.target_url,'Repeated observation','accepted_risk']])
    response = upload(signed_in, assessment, payload)
    assert response.status_code == 302
    batch = db.session.scalar(db.select(ImportBatch))
    assert len(batch.records) == 2 and not batch.errors
    assert db.session.scalar(db.select(Finding)) is None
    preview = signed_in.get(response.location)
    assert b'1</strong> potential duplicates' in preview.data
    signed_in.post(f'/imports/{batch.id}/confirm')
    findings = db.session.scalars(db.select(Finding).order_by(Finding.id)).all()
    assert len(findings) == 2
    assert all(f.status == 'open' and f.confidence == 'Needs review' and f.source == 'Excel import' for f in findings)
    assert findings[1].duplicate_of_id == findings[0].id
    assert batch.status == 'committed' and batch.records == []
    signed_in.post(f'/imports/{batch.id}/confirm')
    assert db.session.scalar(db.select(db.func.count()).select_from(Finding)) == 2


def test_row_errors_block_all_rows_and_formulas_are_rejected(signed_in, assessment):
    payload = workbook_bytes([['Title','Severity'],['Valid row','Low'],['Invalid severity','Urgent'],['=1+2','High']])
    upload(signed_in, assessment, payload)
    batch = db.session.scalar(db.select(ImportBatch))
    assert len(batch.records) == 1 and len(batch.errors) == 2
    assert [row['row'] for row in batch.errors] == [3, 4]
    response = signed_in.get(f'/imports/{batch.id}')
    assert b'Formula cells' in response.data
    assert b'Import 1 findings' not in response.data
    signed_in.post(f'/imports/{batch.id}/confirm')
    assert db.session.scalar(db.select(Finding)) is None


def test_preview_expiration_discard_and_cross_user_access(app, signed_in, assessment):
    upload(signed_in, assessment, workbook_bytes([['Title','Severity'],['Sample','Medium']]))
    batch = db.session.scalar(db.select(ImportBatch))
    other = app.test_client()
    other.post('/login', data={'email':'bob@example.com','password':'TestPassword-Only123!'})
    assert other.get(f'/imports/{batch.id}').status_code == 404
    assert other.post(f'/imports/{batch.id}/confirm').status_code == 404
    assert other.post(f'/imports/{batch.id}/discard').status_code == 404
    batch.created_at = now() - timedelta(minutes=31)
    db.session.commit()
    from vulntrack.services.scanner import process_next
    process_next(app)
    assert batch.status == 'expired' and batch.records == []
    signed_in.post(f'/imports/{batch.id}/confirm')
    assert db.session.scalar(db.select(Finding)) is None
    upload(signed_in, assessment, workbook_bytes([['Title','Severity'],['Another preview','Low']]))
    pending = db.session.scalar(db.select(ImportBatch).where(ImportBatch.status == 'pending'))
    signed_in.post(f'/imports/{pending.id}/discard')
    assert pending.status == 'discarded' and pending.records == []


def test_csv_aliases_and_invalid_numbers():
    rows, errors, warnings = spreadsheet_records(b'Issue Name;Risk Rating;Affected URL;Business Impact\nSample;high;https://example.com;0\nInvalid;low;https://example.com;101\n', 'csv', 'https://fallback.example.com')
    assert len(rows) == 1 and rows[0]['severity'] == 'High' and rows[0]['business_impact'] == 0
    assert len(errors) == 1 and errors[0]['row'] == 3
    assert warnings


def test_missing_headers_invalid_archive_and_oversize_limits():
    with pytest.raises(ValueError, match='Required columns'):
        spreadsheet_records(b'Name,Other\nExample,Value', 'csv', 'https://example.com')
    with pytest.raises(ValueError, match='readable'):
        spreadsheet_records(b'not a zip', 'xlsx', 'https://example.com')
    payload = ('Title,Severity\n' + 'Finding,Low\n' * 2001).encode()
    with pytest.raises(ValueError, match='2,000|2000'):
        spreadsheet_records(payload, 'csv', 'https://example.com')


def test_excel_does_not_trust_truncated_dimensions():
    original = workbook_bytes([['Title','Severity'],['One','Low'],['Two','High']])
    rebuilt = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(rebuilt, 'w') as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                data = re.sub(br'<dimension ref="[^"]+"', b'<dimension ref="A1"', data)
            target.writestr(entry, data)
    records, errors, _ = spreadsheet_records(rebuilt.getvalue(), 'xlsx', 'https://example.com')
    assert len(records) == 2 and not errors


def test_excel_import_real_csrf(app, client, assessment):
    app.config['WTF_CSRF_ENABLED'] = True
    def token(response):
        return re.search(rb'name="csrf_token"[^>]*value="([^"]+)"', response.data)[1].decode()
    csrf = token(client.get('/login'))
    client.post('/login', data={'csrf_token':csrf,'email':'alice@example.com','password':'TestPassword-Only123!'})
    csrf = token(client.get(f'/assessments/{assessment.id}/upload-report'))
    response = client.post(f'/assessments/{assessment.id}/upload-report', data={'csrf_token':csrf,'report':(BytesIO(workbook_bytes([['Title','Severity'],['CSRF fixture','Low']])),'fixture.xlsx'),'scope_confirmed':'y'})
    assert response.status_code == 302
    batch = db.session.scalar(db.select(ImportBatch))
    assert client.post(f'/imports/{batch.id}/confirm').status_code == 400
    csrf = token(client.get(response.location))
    assert client.post(f'/imports/{batch.id}/confirm', data={'csrf_token':csrf}).status_code == 302
    assert db.session.scalar(db.select(Finding))


def test_active_requires_separate_acknowledgement_and_scope(app, signed_in, assessment, monkeypatch):
    from vulntrack.services import scanner, zap_engine
    from vulntrack.services.authorization import active_approval
    app.config['ZAP_ENABLED'] = True
    monkeypatch.setattr(zap_engine, 'available', lambda app: True)
    monkeypatch.setattr(scanner, 'validated_target', lambda target: None)
    assessment.scan_approved = True
    db.session.commit()
    signed_in.post(f'/assessments/{assessment.id}/scan', data={'confirmed':'y','active':'y'})
    assert db.session.scalar(db.select(ScanJob)) is None
    assert signed_in.post(f'/assessments/{assessment.id}/approve-active', data={'confirmed':'y'}).status_code == 403
    signed_in.post('/logout')
    signed_in.post('/login', data={'email':'admin@example.com','password':'TestPassword-Only123!'})
    signed_in.post(f'/assessments/{assessment.id}/approve-active')
    assert active_approval(assessment) is None
    signed_in.post(f'/assessments/{assessment.id}/approve-active', data={'confirmed':'y'})
    assert active_approval(assessment)
    signed_in.post(f'/assessments/{assessment.id}/scan', data={'confirmed':'y'})
    job = db.session.scalar(db.select(ScanJob))
    assert job.result['profile'] == 'passive'
    signed_in.post(f'/scans/{job.id}/cancel')
    signed_in.post(f'/assessments/{assessment.id}/scan', data={'confirmed':'y','active':'y'})
    active_job = db.session.scalar(db.select(ScanJob).where(ScanJob.status == 'queued'))
    assert active_job.result['profile'] == 'active'
    assessment.scope += ' Changed scope.'
    db.session.commit()
    assert active_approval(assessment) is None
    scanner.process_next(app)
    assert active_job.status == 'failed' and 'scope changed' in active_job.error


def test_active_approval_expires(app, assessment):
    from vulntrack.services.authorization import active_approval, scope_digest
    assessment.scan_approved = True
    approval = ActiveScanApproval(assessment_id=assessment.id, approved_by_id=3, scope_digest=scope_digest(assessment), created_at=now()-timedelta(days=8))
    db.session.add(approval)
    db.session.commit()
    assert active_approval(assessment) is None


def test_progress_and_status_are_scoped(app, signed_in, assessment):
    job = ScanJob(assessment_id=assessment.id, requested_by_id=1, target_url=assessment.target_url)
    db.session.add(job)
    db.session.commit()
    assert signed_in.get(f'/scans/{job.id}').status_code == 200
    assert signed_in.get(f'/scans/{job.id}/status').json['status'] == 'queued'
    other = app.test_client()
    other.post('/login', data={'email':'bob@example.com','password':'TestPassword-Only123!'})
    assert other.get(f'/scans/{job.id}').status_code == 404
    assert other.get(f'/scans/{job.id}/status').status_code == 404
