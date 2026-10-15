import hashlib
import json
import re
from io import BytesIO

import pytest
from openpyxl import load_workbook

from vulntrack.extensions import db
from vulntrack.models import Assessment, AuditLog, Finding, ImportBatch, Report, ScanJob, User
from vulntrack.services.demo import seed
from vulntrack.services.intake import add_finding


def finding_data(**updates):
    data = {'title': 'SQL injection in search', 'url': 'https://portal.example.com/search?q=one', 'method': 'GET', 'parameter': 'q', 'cwe': '89',
            'severity': 'High', 'confidence': 'High', 'category': 'auto', 'description': 'Untrusted input was used in a query.',
            'evidence': '<script>alert(1)</script> & sample evidence', 'remediation': 'Use parameterized queries.',
            'exploitability': '80', 'business_impact': '90', 'exposure': '75'}
    return {**data, **updates}


def create_finding(a, **overrides):
    return add_finding(a, {'title': 'Sample finding', 'url': a.target_url, 'severity': 'Medium', 'description': 'Sample impact', 'remediation': 'Apply fix', **overrides})


def test_authentication_and_security_headers(client):
    assert client.get('/').status_code == 302
    assert client.post('/login', data={'email': 'alice@example.com', 'password': 'wrong'}).status_code == 200
    response = client.post('/login', data={'email': 'alice@example.com', 'password': 'TestPassword-Only123!'})
    assert response.status_code == 302
    assert 'HttpOnly' in response.headers['Set-Cookie']
    assert 'SameSite=Lax' in response.headers['Set-Cookie']
    response = client.get('/')
    assert response.status_code == 200
    assert response.headers['X-Frame-Options'] == 'DENY'
    assert response.headers['X-Content-Type-Options'] == 'nosniff'
    assert "script-src 'self'" in response.headers['Content-Security-Policy']
    assert response.headers['Cache-Control'] == 'no-store'
    assert client.post('/logout').status_code == 302
    assert client.get('/').status_code == 302


def test_login_throttled_and_disabled(client):
    for _ in range(10):
        assert client.post('/login', data={'email': 'unknown@example.com', 'password': 'WrongPassword'}).status_code == 200
    assert client.post('/login', data={'email': 'alice@example.com', 'password': 'TestPassword-Only123!'}).status_code == 429


def test_assessment_create_edit_retest_and_owner(signed_in):
    payload = {'title': 'New baseline', 'client': 'Authorized lab', 'target_url': 'https://portal.example.com', 'scope': 'Application assessment excluding destructive actions.', 'authorization_ref': 'LAB-001', 'kind': 'baseline', 'baseline_id': 0, 'authorized': 'y'}
    response = signed_in.post('/assessments/new', data=payload)
    assert response.status_code == 302
    a = db.session.scalar(db.select(Assessment))
    assert a.owner_id == 1
    assert signed_in.get(f'/assessments/{a.id}').status_code == 200
    a.scan_approved = True
    db.session.commit()
    response = signed_in.post(f'/assessments/{a.id}/edit', data={**payload, 'target_url': 'https://other.example.com'})
    assert response.status_code == 302
    db.session.refresh(a)
    assert not a.scan_approved
    response = signed_in.post('/assessments/new', data={**payload, 'title': 'Retest', 'kind': 'retest', 'baseline_id': a.id})
    assert response.status_code == 302
    r = db.session.scalar(db.select(Assessment).where(Assessment.kind == 'retest'))
    assert r.baseline_id == a.id
    response = signed_in.post(f'/assessments/{a.id}/edit', data={**payload, 'kind': 'retest', 'baseline_id': r.id})
    assert response.status_code == 200
    assert a.kind == 'baseline'


def test_finding_lifecycle_duplicate_risk_and_escaping(signed_in, assessment):
    response = signed_in.post(f'/assessments/{assessment.id}/findings/new', data=finding_data())
    assert response.status_code == 302
    f = db.session.scalar(db.select(Finding))
    assert f.category == 'A05:2025'
    assert f.risk_score == 82
    response = signed_in.get(f'/findings/{f.id}')
    assert b'&lt;script&gt;alert(1)&lt;/script&gt;' in response.data
    assert b'<script>alert(1)</script>' not in response.data
    signed_in.post(f'/assessments/{assessment.id}/findings/new', data=finding_data())
    duplicate = db.session.scalar(db.select(Finding).where(Finding.id != f.id))
    assert duplicate.duplicate_of_id == f.id
    signed_in.post(f'/findings/{duplicate.id}/duplicate-dismiss')
    db.session.refresh(duplicate)
    assert duplicate.duplicate_of_id is None
    signed_in.post(f'/findings/{f.id}/status', data={'status': 'accepted_risk', 'resolution_note': ''})
    db.session.refresh(f)
    assert f.status == 'open'
    signed_in.post(f'/findings/{f.id}/status', data={'status': 'remediated', 'resolution_note': 'Retest confirms input is parameterized.'})
    assert f.status == 'open'
    signed_in.post(f'/findings/{f.id}/status', data={'status': 'remediated', 'resolution_note': 'Retest confirms input is parameterized.', 'retest_confirmed': 'y', 'assigned_to': 'App team'})
    db.session.refresh(f)
    assert f.status == 'remediated' and f.retested_by_id == 1 and f.retested_at
    signed_in.post(f'/findings/{f.id}/status', data={'status': 'open', 'resolution_note': 'Regression observed in authorized retest.'})
    db.session.refresh(f)
    assert f.status == 'open' and f.retested_at is None


@pytest.mark.parametrize('url', ['file:///etc/passwd', 'javascript:alert(1)', 'https://user:pass@example.com', 'https://example.com\r\nX-Test:yes', 'http://example.com:99999'])
def test_invalid_finding_urls_rejected(signed_in, assessment, url):
    response = signed_in.post(f'/assessments/{assessment.id}/findings/new', data=finding_data(url=url))
    assert response.status_code == 200
    assert db.session.scalar(db.select(db.func.count()).select_from(Finding)) == 0


def test_reports_snapshot_excel_formula_safety_and_pdf(signed_in, assessment):
    title = '=HYPERLINK("https://attacker.invalid","click")'
    f = create_finding(assessment, title=title, evidence='<b>Evidence</b> & details\n' + 'Long evidence. ' * 500)
    db.session.commit()
    for fmt in ['xlsx', 'pdf']:
        response = signed_in.post(f'/assessments/{assessment.id}/reports', data={'format': fmt})
        assert response.status_code == 302
        report = db.session.scalar(db.select(Report).where(Report.format == fmt))
        response = signed_in.get(f'/reports/{report.id}/download')
        assert response.status_code == 200
        assert hashlib.sha256(response.data).hexdigest() == report.sha256
        assert 'attachment' in response.headers['Content-Disposition']
        if fmt == 'xlsx':
            wb = load_workbook(BytesIO(response.data))
            assert wb.sheetnames == ['Summary', 'Findings', 'Remediation']
            assert wb['Findings']['B2'].value == title
            assert wb['Findings']['B2'].data_type == 's'
            assert not any(c.data_type == 'f' for sheet in wb for row in sheet for c in row)
        else:
            assert response.data.startswith(b'%PDF-')
        old_content = report.content
        f.title = 'Edited after generation'
        db.session.commit()
        db.session.refresh(report)
        assert report.content == old_content
        assert report.snapshot['findings'][0]['title'] == title
        f.title = title
        db.session.commit()
    assert db.session.scalar(db.select(AuditLog).where(AuditLog.action == 'report_downloaded'))


def test_user_cannot_access_other_users_records_or_admin(signed_in, assessment):
    f = create_finding(assessment)
    db.session.commit()
    signed_in.post(f'/assessments/{assessment.id}/reports', data={'format': 'xlsx'})
    report = db.session.scalar(db.select(Report))
    job = ScanJob(assessment_id=assessment.id, requested_by_id=1, target_url=assessment.target_url)
    db.session.add(job)
    db.session.commit()
    signed_in.post('/logout')
    signed_in.post('/login', data={'email': 'bob@example.com', 'password': 'TestPassword-Only123!'})
    for path in [f'/assessments/{assessment.id}', f'/assessments/{assessment.id}/edit', f'/assessments/{assessment.id}/import', f'/findings/{f.id}', f'/findings/{f.id}/edit', f'/reports/{report.id}/download', f'/?assessment_id={assessment.id}', f'/findings?assessment_id={assessment.id}']:
        assert signed_in.get(path).status_code == 404, path
    for path, payload in [(f'/findings/{f.id}/status', {'status': 'remediated'}), (f'/scans/{job.id}/cancel', {}), (f'/assessments/{assessment.id}/reports', {'format': 'pdf'}), (f'/assessments/{assessment.id}/scan', {'confirmed': 'y'})]:
        assert signed_in.post(path, data=payload).status_code == 404, path
    for path in ['/users', '/audit']:
        assert signed_in.get(path).status_code == 403
    assert signed_in.post(f'/assessments/{assessment.id}/approve').status_code == 403
    assert b'Authorized portal baseline' not in signed_in.get('/assessments').data
    assert b'Sample finding' not in signed_in.get('/findings').data
    assert b'Snapshot #' not in signed_in.get('/reports').data


def test_password_change_and_disable_invalidate_sessions(app, signed_in):
    second = app.test_client()
    second.post('/login', data={'email': 'alice@example.com', 'password': 'TestPassword-Only123!'})
    signed_in.post('/account', data={'current_password': 'TestPassword-Only123!', 'password': 'NewLongPassword123!', 'confirm': 'NewLongPassword123!'})
    assert signed_in.get('/').status_code == 302
    assert second.get('/').status_code == 302
    signed_in.post('/login', data={'email': 'alice@example.com', 'password': 'NewLongPassword123!'})
    admin = app.test_client()
    admin.post('/login', data={'email': 'admin@example.com', 'password': 'TestPassword-Only123!'})
    assert admin.post('/users/1/toggle').status_code == 302
    assert signed_in.get('/').status_code == 302
    assert admin.post('/users/3/toggle').status_code == 302
    assert db.session.get(User, 3).enabled


def test_real_csrf_login_and_mutations(app, client, assessment):
    app.config['WTF_CSRF_ENABLED'] = True
    assert client.post('/login', data={'email': 'alice@example.com', 'password': 'TestPassword-Only123!'}).status_code == 400
    response = client.get('/login')
    token = re.search(rb'name="csrf_token"[^>]*value="([^"]+)"', response.data).group(1).decode()
    assert client.post('/login', data={'csrf_token': token, 'email': 'alice@example.com', 'password': 'TestPassword-Only123!'}).status_code == 302
    assert client.post(f'/assessments/{assessment.id}/reports', data={'format': 'pdf'}).status_code == 400
    response = client.get(f'/assessments/{assessment.id}')
    token = re.search(rb'name="csrf_token"[^>]*value="([^"]+)"', response.data).group(1).decode()
    assert client.post(f'/assessments/{assessment.id}/reports', data={'csrf_token': token, 'format': 'pdf'}).status_code == 302
    assert client.post('/logout', data={'csrf_token': token}).status_code == 302


def test_all_pages_seeded_demo_and_pagination(app, signed_in):
    assert seed(db.session.get(User, 1))
    assert not seed(db.session.get(User, 1))
    assert db.session.scalar(db.select(db.func.count()).select_from(Finding)) == 25
    a = db.session.scalar(db.select(Assessment).where(Assessment.kind == 'baseline'))
    retest = db.session.scalar(db.select(Assessment).where(Assessment.kind == 'retest'))
    paths = ['/', '/assessments', '/assessments/new', f'/assessments/{a.id}', f'/assessments/{a.id}/edit', f'/assessments/{a.id}/findings/new', f'/assessments/{a.id}/import',
             '/findings', '/findings?page=2', '/findings?severity=High&sort=newest', '/remediation', '/reports', '/scans', '/account', '/comparison', f'/comparison?baseline_id={a.id}&retest_id={retest.id}']
    for path in paths:
        response = signed_in.get(path)
        assert response.status_code == 200, path
    signed_in.post('/logout')
    signed_in.post('/login', data={'email': 'admin@example.com', 'password': 'TestPassword-Only123!'})
    for path in ['/users', '/audit', '/', '/reports']:
        assert signed_in.get(path).status_code == 200


def test_import_atomic_and_parsing_errors(signed_in, assessment):
    report = {'site': [{'@name': assessment.target_url, 'alerts': [{'alert': 'Cookie flag missing', 'riskcode': '1', 'cweid': '614', 'desc': '<p>Description</p>', 'solution': '<p>Set flags</p>', 'instances': [{'uri': assessment.target_url, 'method': 'GET', 'evidence': '<b>Cookie</b>'}]}]}]}
    result = signed_in.post(f'/assessments/{assessment.id}/import', data={'report': (BytesIO(json.dumps(report).encode()), 'report.json'), 'scope_confirmed': 'y'})
    assert result.status_code == 302
    batch = db.session.scalar(db.select(ImportBatch))
    assert batch and db.session.scalar(db.select(db.func.count()).select_from(Finding)) == 0
    assert signed_in.get(result.location).status_code == 200
    assert signed_in.post(f'/imports/{batch.id}/confirm').status_code == 302
    f = db.session.scalar(db.select(Finding))
    assert f.description == 'Description' and f.evidence == 'Cookie'
    report['site'][0]['alerts'].append({'alert': 'Invalid second finding', 'riskcode': '3', 'instances': [{'uri': 'file:///etc/passwd'}]})
    signed_in.post(f'/assessments/{assessment.id}/import', data={'report': (BytesIO(json.dumps(report).encode()), 'report.json'), 'scope_confirmed': 'y'})
    assert db.session.scalar(db.select(db.func.count()).select_from(Finding)) == 1
    response = signed_in.post(f'/assessments/{assessment.id}/import', data={'report': (BytesIO(b'<!DOCTYPE foo [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><OWASPZAPReport>&xxe;</OWASPZAPReport>'), 'report.xml'), 'scope_confirmed': 'y'})
    assert db.session.scalar(db.select(db.func.count()).select_from(Finding)) == 1
    assert b'Unsafe or malformed XML' in response.data


def test_content_size_limit(app, signed_in, assessment):
    app.config['MAX_CONTENT_LENGTH'] = 1000
    response = signed_in.post(f'/assessments/{assessment.id}/import', data={'report': (BytesIO(b'x' * 2000), 'large.json')})
    assert response.status_code == 413


def test_delete_reauthentication_dependencies_and_history(signed_in, assessment):
    f = create_finding(assessment)
    duplicate = create_finding(assessment)
    db.session.commit()
    signed_in.post(f'/assessments/{assessment.id}/reports', data={'format': 'xlsx'})
    finding_id = f.id
    assert signed_in.get(f'/findings/{finding_id}/delete').status_code == 200
    assert signed_in.post(f'/findings/{finding_id}/delete', data={'confirmation': 'Wrong title', 'current_password': 'TestPassword-Only123!'}).status_code == 200
    assert db.session.get(Finding, finding_id)
    assert signed_in.post(f'/findings/{finding_id}/delete', data={'confirmation': f.title, 'current_password': 'TestPassword-Only123!'}).status_code == 302
    db.session.refresh(duplicate)
    assert duplicate.duplicate_of_id is None
    assert db.session.scalar(db.select(Report)) is not None
    retest = Assessment(owner_id=1, title='Linked retest', client=assessment.client, target_url=assessment.target_url, scope=assessment.scope, authorization_ref=assessment.authorization_ref, kind='retest', baseline_id=assessment.id)
    db.session.add(retest)
    db.session.commit()
    a_id, r_id = assessment.id, retest.id
    assert signed_in.post(f'/assessments/{a_id}/delete', data={'confirmation': assessment.title, 'current_password': 'TestPassword-Only123!'}).status_code == 200
    assert db.session.get(Assessment, a_id)
    assert signed_in.post(f'/assessments/{r_id}/delete', data={'confirmation': retest.title, 'current_password': 'TestPassword-Only123!'}).status_code == 302
    assert signed_in.post(f'/assessments/{a_id}/delete', data={'confirmation': assessment.title, 'current_password': 'TestPassword-Only123!'}).status_code == 302
    assert db.session.scalar(db.select(Report)) is None
    assert db.session.scalar(db.select(Finding)) is None
    assert db.session.scalar(db.select(AuditLog).where(AuditLog.action == 'assessment_deleted'))


def test_admin_create_account_validation_and_permissions(app, client):
    client.post('/login', data={'email': 'admin@example.com', 'password': 'TestPassword-Only123!'})
    response = client.post('/users', data={'name': 'New Tester', 'email': 'new@example.com', 'password': 'LongNewPassword123!', 'role': 'tester'})
    assert response.status_code == 302
    user = db.session.scalar(db.select(User).where(User.email == 'new@example.com'))
    assert user and user.password_hash != 'LongNewPassword123!' and user.check_password('LongNewPassword123!')
    response = client.post('/users', data={'name': 'Duplicate', 'email': 'new@example.com', 'password': 'LongNewPassword123!', 'role': 'admin'})
    assert response.status_code == 200
    assert user.role == 'tester'


def test_report_generation_limit(signed_in, assessment, monkeypatch):
    from vulntrack.services import reports
    monkeypatch.setattr(reports, 'snapshot', lambda *args: (_ for _ in ()).throw(ValueError('A report is limited to 500 findings.')))
    response = signed_in.post(f'/assessments/{assessment.id}/reports', data={'format': 'pdf'}, follow_redirects=True)
    assert b'limited to 500' in response.data
    assert db.session.scalar(db.select(Report)) is None
