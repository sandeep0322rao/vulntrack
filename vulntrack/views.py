import hashlib
import secrets
from collections import Counter
from datetime import timedelta
from functools import wraps
from io import BytesIO

from flask import Blueprint, abort, current_app, flash, redirect, render_template, request, send_file, session, url_for
from flask_login import current_user, login_required, login_user, logout_user
from sqlalchemy import or_
from sqlalchemy.exc import IntegrityError
from werkzeug.security import check_password_hash, generate_password_hash

from .constants import CATEGORIES, OPEN_STATUSES, SEVERITIES, STATUSES
from .extensions import db
from .forms import ActionForm, ActiveApprovalForm, AssessmentForm, DeleteForm, FindingForm, ImportForm, LoginForm, PasswordForm, ReportForm, ScanForm, StatusForm, UserForm
from .models import ActiveScanApproval, Assessment, AuditLog, Finding, ImportBatch, LoginAttempt, Report, ScanJob, User, now
from .services.intake import add_finding, zap_records
from .services.risk import fingerprint, score, suggest_category

web = Blueprint('web', __name__)
DUMMY_HASH = generate_password_hash(secrets.token_urlsafe(24), method='scrypt')


def admin_required(func):
    @wraps(func)
    @login_required
    def wrapped(*args, **kwargs):
        if not current_user.is_admin:
            abort(403)
        return func(*args, **kwargs)
    return wrapped


def assessments_query():
    query = db.select(Assessment)
    return query if current_user.is_admin else query.where(Assessment.owner_id == current_user.id)


def findings_query():
    query = db.select(Finding).join(Assessment)
    return query if current_user.is_admin else query.where(Assessment.owner_id == current_user.id)


def get_assessment(identifier):
    return db.session.scalar(assessments_query().where(Assessment.id == identifier)) or abort(404)


def get_finding(identifier):
    return db.session.scalar(findings_query().where(Finding.id == identifier)) or abort(404)


def audit(action, assessment=None, details=''):
    db.session.add(AuditLog(actor_id=current_user.id if current_user.is_authenticated else None,
                            assessment_id=assessment.id if assessment else None, action=action, details=details[:1000]))


def form_errors(form):
    for name, errors in form.errors.items():
        for error in errors:
            label = getattr(form, name).label.text
            flash(f'{label}: {error}', 'danger')


def assessment_choices(exclude=None):
    query = assessments_query().where(Assessment.kind == 'baseline')
    if exclude:
        query = query.where(Assessment.id != exclude)
    return [(0, 'Select a baseline (retests only)')] + [(a.id, a.title) for a in db.session.scalars(query.order_by(Assessment.created_at.desc()))]


def selected_assessment():
    raw = request.args.get('assessment_id', '').strip()
    if not raw:
        return None
    if not raw.isdigit():
        abort(400)
    return get_assessment(int(raw))


@web.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('web.dashboard'))
    form = LoginForm()
    if form.validate_on_submit():
        ip_key = hashlib.sha256((request.remote_addr or 'unknown').encode()).hexdigest()
        cutoff = now() - timedelta(minutes=10)
        attempts = db.session.scalar(db.select(db.func.count()).select_from(LoginAttempt).where(LoginAttempt.ip_key == ip_key, LoginAttempt.created_at >= cutoff))
        if attempts >= 10:
            abort(429)
        db.session.execute(db.delete(LoginAttempt).where(LoginAttempt.created_at < cutoff))
        user = db.session.scalar(db.select(User).where(User.email == form.email.data.strip().lower()))
        valid = check_password_hash(user.password_hash if user else DUMMY_HASH, form.password.data)
        if not user or not valid or not user.enabled:
            db.session.add(LoginAttempt(ip_key=ip_key))
            db.session.add(AuditLog(action='login_failed', details='Sign-in rejected'))
            db.session.commit()
            flash('Email or password is incorrect, or the account is disabled.', 'danger')
        else:
            session.clear()
            login_user(user, remember=False)
            session.permanent = True
            audit('login')
            db.session.commit()
            return redirect(url_for('web.dashboard'))
    return render_template('login.html', form=form)


@web.post('/logout')
@login_required
def logout():
    audit('logout')
    db.session.commit()
    logout_user()
    session.clear()
    return redirect(url_for('web.login'))


@web.get('/')
@login_required
def dashboard():
    assessment = selected_assessment()
    query = findings_query()
    if assessment:
        query = query.where(Finding.assessment_id == assessment.id)
    findings = list(db.session.scalars(query))
    assessments = list(db.session.scalars(assessments_query().order_by(Assessment.created_at.desc())))
    metrics = {
        'assessments': 1 if assessment else len(assessments), 'total': len(findings),
        'open': sum(f.status in OPEN_STATUSES for f in findings),
        'urgent': sum(f.status in OPEN_STATUSES and f.risk_score >= 60 for f in findings),
        'remediated': sum(f.status == 'remediated' for f in findings),
        'retest': sum(f.status == 'ready_retest' for f in findings),
    }
    from .services.risk import priority
    charts = {
        'severity': {s: sum(f.severity == s for f in findings) for s in SEVERITIES},
        'status': {label: sum(f.status == key for f in findings) for key, label in STATUSES.items()},
        'category': {k: sum(f.category == k for f in findings) for k in CATEGORIES if any(f.category == k for f in findings)},
        'priority': {s: sum(priority(f.risk_score) == s for f in findings) for s in ['Critical', 'High', 'Medium', 'Low']},
    }
    recent = sorted((f for f in findings if f.status in OPEN_STATUSES), key=lambda f: (-f.risk_score, -f.id))[:5]
    jobs_query = db.select(ScanJob).join(Assessment)
    if not current_user.is_admin:
        jobs_query = jobs_query.where(Assessment.owner_id == current_user.id)
    if assessment:
        jobs_query = jobs_query.where(Assessment.id == assessment.id)
    jobs = list(db.session.scalars(jobs_query.order_by(ScanJob.id.desc()).limit(4)))
    demo_scope = ('Synthetic' in assessment.client) if assessment else bool(assessments and all('Synthetic' in a.client for a in assessments))
    return render_template('dashboard.html', assessment=assessment, assessments=assessments, metrics=metrics, charts=charts, recent=recent, jobs=jobs, demo_scope=demo_scope)


@web.get('/assessments')
@login_required
def assessments():
    search = request.args.get('q', '').strip()[:160]
    query = assessments_query()
    if search:
        query = query.where(or_(Assessment.title.contains(search, autoescape=True), Assessment.client.contains(search, autoescape=True)))
    pagination = db.paginate(query.order_by(Assessment.created_at.desc()), per_page=12, max_per_page=12)
    return render_template('assessments.html', pagination=pagination, search=search)


@web.route('/assessments/new', methods=['GET', 'POST'])
@login_required
def assessment_new():
    form = AssessmentForm()
    form.baseline_id.choices = assessment_choices()
    if form.validate_on_submit():
        baseline = get_assessment(form.baseline_id.data) if form.kind.data == 'retest' and form.baseline_id.data else None
        if form.kind.data == 'retest' and not baseline:
            form.baseline_id.errors.append('Select a baseline assessment for this retest.')
        elif baseline and (baseline.owner_id != current_user.id or baseline.kind != 'baseline'):
            form.baseline_id.errors.append('Choose your own baseline assessment.')
        else:
            a = Assessment(owner_id=current_user.id)
            for field in ['title', 'client', 'target_url', 'scope', 'authorization_ref', 'kind']:
                setattr(a, field, getattr(form, field).data.strip())
            a.baseline_id = baseline.id if baseline else None
            db.session.add(a)
            db.session.flush()
            audit('assessment_created', a, a.title)
            db.session.commit()
            flash('Assessment created. Choose how to add your findings.', 'success')
            return redirect(url_for('web.intake', identifier=a.id))
    return render_template('assessment_form.html', form=form, assessment=None)


@web.route('/assessments/<int:identifier>/edit', methods=['GET', 'POST'])
@login_required
def assessment_edit(identifier):
    a = get_assessment(identifier)
    form = AssessmentForm(obj=a)
    form.baseline_id.choices = assessment_choices(exclude=a.id)
    if request.method == 'GET':
        form.authorized.data = True
        form.baseline_id.data = a.baseline_id or 0
    if form.validate_on_submit():
        baseline = get_assessment(form.baseline_id.data) if form.kind.data == 'retest' and form.baseline_id.data else None
        active = db.session.scalar(db.select(ScanJob).where(ScanJob.assessment_id == a.id, ScanJob.status.in_(['queued', 'running'])))
        has_retests = db.session.scalar(db.select(Assessment).where(Assessment.baseline_id == a.id))
        if active:
            flash('Cancel or finish the active scan before editing its scope.', 'warning')
        elif has_retests and form.kind.data != 'baseline':
            flash('This baseline has linked retests and cannot be converted to a retest.', 'warning')
        elif form.kind.data == 'retest' and (not baseline or baseline.owner_id != a.owner_id or baseline.kind != 'baseline'):
            form.baseline_id.errors.append('Choose a baseline belonging to the same owner.')
        else:
            # Any scope or authorization change invalidates previous scan approval.
            scope_changed = any(getattr(a, key) != getattr(form, key).data.strip() for key in ['target_url', 'scope', 'authorization_ref'])
            for key in ['title', 'client', 'target_url', 'scope', 'authorization_ref', 'kind']:
                setattr(a, key, getattr(form, key).data.strip())
            a.baseline_id = baseline.id if baseline else None
            if scope_changed:
                a.scan_approved = False
                a.approved_at = None
                a.approved_by_id = None
            audit('assessment_updated', a, 'Scan approval reset' if scope_changed else 'Details updated')
            db.session.commit()
            flash('Assessment updated.', 'success')
            return redirect(url_for('web.assessment_detail', identifier=a.id))
    return render_template('assessment_form.html', form=form, assessment=a)


@web.get('/assessments/<int:identifier>')
@login_required
def assessment_detail(identifier):
    a = get_assessment(identifier)
    recent = db.session.scalars(db.select(Finding).where(Finding.assessment_id == a.id).order_by(Finding.risk_score.desc()).limit(8)).all()
    logs = db.session.scalars(db.select(AuditLog).where(AuditLog.assessment_id == a.id).order_by(AuditLog.id.desc()).limit(8)).all()
    jobs = db.session.scalars(db.select(ScanJob).where(ScanJob.assessment_id == a.id).order_by(ScanJob.id.desc()).limit(5)).all()
    return render_template('assessment_detail.html', assessment=a, recent=recent, logs=logs, jobs=jobs, report_form=ReportForm(), scan_form=ScanForm())


@web.get('/assessments/<int:identifier>/add-findings')
@login_required
def intake(identifier):
    assessment = get_assessment(identifier)
    return render_template('intake.html', assessment=assessment)


@web.get('/assessments/<int:identifier>/automatic-scan')
@login_required
def automatic_scan(identifier):
    assessment = get_assessment(identifier)
    from .services.zap_engine import available
    from .services.authorization import active_approval
    return render_template('automatic_scan.html', assessment=assessment, form=ScanForm(), zap_ready=available(current_app),
                           active_approval=active_approval(assessment), active_form=ActiveApprovalForm())


@web.post('/assessments/<int:identifier>/approve-active')
@admin_required
def active_scan_approve(identifier):
    assessment = get_assessment(identifier)
    form = ActiveApprovalForm()
    if not form.validate_on_submit() or not assessment.scan_approved:
        flash('Approve the target first, then explicitly confirm authorization for active tests.', 'warning')
    else:
        from .services.scanner import validated_target
        from .services.authorization import scope_digest
        try:
            validated_target(assessment.target_url)
        except ValueError as exc:
            flash(str(exc), 'danger')
        else:
            approval = db.session.scalar(db.select(ActiveScanApproval).where(ActiveScanApproval.assessment_id == assessment.id))
            if not approval:
                approval = ActiveScanApproval(assessment_id=assessment.id)
                db.session.add(approval)
            approval.approved_by_id, approval.scope_digest, approval.created_at = current_user.id, scope_digest(assessment), now()
            audit('active_scan_scope_approved', assessment, 'SQL injection and reflected XSS testing permitted for seven days; explicit opt-in required per scan')
            db.session.commit()
            flash('Active testing approved. Low-impact scanning remains the default.', 'success')
    return redirect(url_for('web.automatic_scan', identifier=assessment.id))


@web.post('/assessments/<int:identifier>/approve')
@admin_required
def assessment_approve(identifier):
    a = get_assessment(identifier)
    from .services.scanner import validated_target
    try:
        validated_target(a.target_url)
    except ValueError as exc:
        flash(f'Scan approval rejected: {exc}', 'danger')
    else:
        a.scan_approved = True
        a.approved_by_id = current_user.id
        a.approved_at = now()
        audit('scan_scope_approved', a, a.target_url)
        db.session.commit()
        flash('Public target approved for bounded passive checks.', 'success')
    return redirect(url_for('web.automatic_scan', identifier=a.id))


@web.get('/findings')
@login_required
def findings():
    assessment = selected_assessment()
    query = findings_query()
    if assessment:
        query = query.where(Finding.assessment_id == assessment.id)
    filters = {k: request.args.get(k, '').strip()[:240] for k in ['q', 'severity', 'status', 'category', 'sort']}
    if filters['q']:
        query = query.where(or_(Finding.title.contains(filters['q'], autoescape=True), Finding.url.contains(filters['q'], autoescape=True)))
    for key, allowed in [('severity', SEVERITIES), ('status', STATUSES), ('category', CATEGORIES)]:
        if filters[key]:
            if filters[key] not in allowed:
                abort(400)
            query = query.where(getattr(Finding, key) == filters[key])
    ordering = {'risk': Finding.risk_score.desc(), 'newest': Finding.id.desc(), 'title': Finding.title.asc()}
    if filters['sort'] and filters['sort'] not in ordering:
        abort(400)
    pagination = db.paginate(query.order_by(ordering.get(filters['sort'], Finding.risk_score.desc()), Finding.id.desc()), per_page=20, max_per_page=20)
    return render_template('findings.html', assessment=assessment, assessments=db.session.scalars(assessments_query()).all(), pagination=pagination, filters=filters)


def finding_form_data(form):
    fields = ['title', 'url', 'method', 'parameter', 'cwe', 'severity', 'confidence', 'description', 'evidence', 'remediation', 'exploitability', 'business_impact', 'exposure']
    data = {k: getattr(form, k).data.strip() if isinstance(getattr(form, k).data, str) else getattr(form, k).data for k in fields}
    if form.category.data != 'auto':
        data['category'] = form.category.data
        data['mapping_reason'] = 'Classification explicitly selected by assessor.'
    return data


@web.route('/assessments/<int:identifier>/findings/new', methods=['GET', 'POST'])
@login_required
def finding_new(identifier):
    a = get_assessment(identifier)
    form = FindingForm()
    if request.method == 'GET':
        form.url.data = a.target_url
        form.severity.data = 'Medium'
    if form.validate_on_submit():
        f = add_finding(a, finding_form_data(form))
        audit('finding_created', a, f'Finding #{f.id}')
        db.session.commit()
        flash('Finding saved. Review its classification and priority.', 'success')
        return redirect(url_for('web.findings', assessment_id=a.id))
    return render_template('finding_form.html', assessment=a, form=form, finding=None)


@web.route('/findings/<int:identifier>/edit', methods=['GET', 'POST'])
@login_required
def finding_edit(identifier):
    f = get_finding(identifier)
    form = FindingForm(obj=f)
    if form.validate_on_submit():
        data = finding_form_data(form)
        if 'category' not in data:
            data['category'], data['mapping_reason'] = suggest_category(data['title'], data['cwe'], data['description'])
        previous_fp = f.fingerprint
        for key, value in data.items():
            setattr(f, key, value)
        f.risk_score = score(f.severity, f.exploitability, f.business_impact, f.exposure)
        f.fingerprint = fingerprint(f.title, f.url, f.method, f.parameter, f.cwe)
        match = db.session.scalar(db.select(Finding).where(Finding.assessment_id == f.assessment_id, Finding.id != f.id, Finding.fingerprint == f.fingerprint).order_by(Finding.id))
        # Only point at an older record, so links cannot form cycles.
        f.duplicate_of_id = match.id if match and match.id < f.id else None
        if f.fingerprint != previous_fp:
            for dependent in db.session.scalars(db.select(Finding).where(Finding.duplicate_of_id == f.id)):
                dependent.duplicate_of_id = None
        audit('finding_updated', f.assessment, f'Finding #{f.id}')
        db.session.commit()
        flash('Finding updated. Existing report snapshots are unchanged.', 'success')
        return redirect(url_for('web.finding_detail', identifier=f.id))
    return render_template('finding_form.html', assessment=f.assessment, form=form, finding=f)


@web.get('/findings/<int:identifier>')
@login_required
def finding_detail(identifier):
    f = get_finding(identifier)
    return render_template('finding_detail.html', finding=f, form=StatusForm(obj=f))


@web.post('/findings/<int:identifier>/status')
@login_required
def finding_status(identifier):
    f = get_finding(identifier)
    form = StatusForm()
    if form.validate_on_submit():
        note = (form.resolution_note.data or '').strip()
        if form.status.data in ('accepted_risk', 'false_positive', 'remediated') and len(note) < 10:
            flash('Record at least 10 characters of evidence or justification for this status.', 'danger')
        elif form.status.data == 'remediated' and not form.retest_confirmed.data:
            flash('Confirm a completed retest before marking a finding remediated.', 'danger')
        else:
            previous = f.status
            f.status = form.status.data
            f.assigned_to = (form.assigned_to.data or '').strip()
            f.resolution_note = note
            if f.status == 'remediated':
                f.retested_by_id, f.retested_at = current_user.id, now()
            else:
                f.retested_by_id, f.retested_at = None, None
            audit('finding_status_changed', f.assessment, f'Finding #{f.id}: {previous} → {f.status}')
            db.session.commit()
            flash('Remediation record updated.', 'success')
    else:
        form_errors(form)
    return redirect(url_for('web.finding_detail', identifier=f.id))


@web.post('/findings/<int:identifier>/duplicate-dismiss')
@login_required
def duplicate_dismiss(identifier):
    f = get_finding(identifier)
    f.duplicate_of_id = None
    audit('duplicate_link_dismissed', f.assessment, f'Finding #{f.id}; retained as a separate record')
    db.session.commit()
    flash('Duplicate link dismissed; both findings are retained.', 'success')
    return redirect(url_for('web.finding_detail', identifier=f.id))


@web.route('/assessments/<int:identifier>/import', methods=['GET', 'POST'])
@web.route('/assessments/<int:identifier>/upload-report', methods=['GET', 'POST'])
@login_required
def import_zap(identifier):
    a = get_assessment(identifier)
    form = ImportForm()
    if form.validate_on_submit():
        upload = form.report.data
        try:
            extension = upload.filename.rsplit('.', 1)[-1].lower()
            payload = upload.read()
            errors, warnings = [], []
            if extension in ('xlsx', 'csv'):
                from .services.spreadsheet_import import spreadsheet_records
                records, errors, warnings = spreadsheet_records(payload, extension, a.target_url)
            else:
                records = zap_records(payload, extension)
                for number, record in enumerate(records, 1):
                    record['_row'] = number
                warnings.append('Scanner alerts need assessor review. Business context and risk inputs use the documented defaults.')
            from werkzeug.utils import secure_filename
            batch = ImportBatch(assessment_id=a.id, owner_id=current_user.id, filename=secure_filename(upload.filename)[:160] or 'report.' + extension,
                                format=extension, records=records, errors=errors, warnings=warnings)
            db.session.add(batch)
            audit('report_upload_previewed', a, f'{len(records)} valid rows; {len(errors)} invalid rows; {extension}')
            db.session.commit()
            return redirect(url_for('web.import_preview', batch_id=batch.id))
        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), 'danger')
    return render_template('import.html', assessment=a, form=form)


def get_import_batch(batch_id):
    batch = db.session.get(ImportBatch, batch_id) or abort(404)
    get_assessment(batch.assessment_id)
    if batch.owner_id != current_user.id:
        abort(404)
    return batch


@web.get('/imports/<batch_id>')
@login_required
def import_preview(batch_id):
    batch = get_import_batch(batch_id)
    from .services.risk import priority
    expired = batch.created_at < now() - timedelta(minutes=30)
    seen = set(db.session.scalars(db.select(Finding.fingerprint).where(Finding.assessment_id == batch.assessment_id)))
    duplicates = 0
    preview = []
    for record in batch.records:
        fp = fingerprint(record['title'], record['url'], record.get('method', 'GET'), record.get('parameter', ''), record.get('cwe', ''))
        duplicate = fp in seen
        duplicates += duplicate
        seen.add(fp)
        if len(preview) < 50:
            preview.append(dict(record, duplicate=duplicate))
    return render_template('import_preview.html', assessment=batch.assessment, batch=batch, preview=preview,
                           duplicates=duplicates, expired=expired)


@web.post('/imports/<batch_id>/confirm')
@login_required
def import_confirm(batch_id):
    batch = get_import_batch(batch_id)
    if batch.status != 'pending':
        flash('This upload was already imported or discarded. No rows were added again.', 'info')
    elif batch.created_at < now() - timedelta(minutes=30):
        flash('This preview expired. Upload the report again before importing.', 'warning')
    elif batch.errors or not batch.records:
        flash('Correct every invalid row and upload the file again. Nothing has been imported.', 'danger')
        return redirect(url_for('web.import_preview', batch_id=batch.id))
    else:
        claimed = db.session.execute(db.update(ImportBatch).where(ImportBatch.id == batch.id, ImportBatch.status == 'pending').values(status='committing'))
        if claimed.rowcount:
            count = len(batch.records)
            for record in batch.records:
                add_finding(batch.assessment, {k: v for k, v in record.items() if not k.startswith('_')})
            batch.status, batch.committed_at = 'committed', now()
            batch.records = []
            audit('report_import_confirmed', batch.assessment, f'{count} findings from upload {batch.id} ({batch.format})')
            db.session.commit()
            flash(f'Imported {count} findings. Review their severity, classification, and remediation below.', 'success')
        else:
            db.session.rollback()
            flash('This report has already been processed.', 'info')
    return redirect(url_for('web.findings', assessment_id=batch.assessment_id))


@web.post('/imports/<batch_id>/discard')
@login_required
def import_discard(batch_id):
    batch = get_import_batch(batch_id)
    claimed = db.session.execute(db.update(ImportBatch).where(ImportBatch.id == batch.id, ImportBatch.status == 'pending').values(status='discarded', records=[]))
    if claimed.rowcount:
        audit('report_upload_discarded', batch.assessment, f'Upload {batch.id}')
    db.session.commit()
    return redirect(url_for('web.import_zap', identifier=batch.assessment_id))


@web.get('/templates/findings.xlsx')
@login_required
def import_template():
    from .services.spreadsheet_import import HEADERS
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Findings'
    sheet.append(HEADERS)
    sheet.freeze_panes = 'A2'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='173D39')
        sheet.column_dimensions[cell.column_letter].width = 28
    sheet.column_dimensions['A'].width = 40
    choices = DataValidation(type='list', formula1='"Critical,High,Medium,Low,Informational"')
    choices.errorTitle, choices.error = 'Invalid severity', 'Choose a severity from the list.'
    choices.showErrorMessage = True
    sheet.add_data_validation(choices)
    choices.add('B2:B2001')
    guide = workbook.create_sheet('Instructions')
    guide.append(['Findings import guide'])
    for text in ['Fill the Findings worksheet, one issue per row. Title and Severity are required.',
                 'URL may be blank to use the assessment target. Use HTTP or HTTPS URLs.',
                 'Optional columns: Description, Evidence, Remediation, Method, Parameter, CWE, Category.',
                 'Category should use OWASP 2025 identifiers, for example A05:2025. Blank categories are suggested automatically.',
                 'Exploitability, Business Impact, and Exposure are whole numbers from 0 to 100; defaults are 50, 50, 75.',
                 'Do not include formulas, macros, credentials, or unrelated personal information.',
                 'Upload previews never save findings until you confirm. Correct all row errors before importing.']:
        guide.append([text])
    guide.column_dimensions['A'].width = 120
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name='vulntrack-findings-template.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@web.get('/remediation')
@login_required
def remediation():
    assessment = selected_assessment()
    query = findings_query()
    if assessment:
        query = query.where(Finding.assessment_id == assessment.id)
    findings = db.session.scalars(query.order_by(Finding.risk_score.desc()).limit(300)).all()
    columns = {key: [f for f in findings if f.status == key] for key in STATUSES}
    return render_template('remediation.html', assessment=assessment, columns=columns,
                           assessments=db.session.scalars(assessments_query()).all(), truncated=len(findings) == 300)


@web.get('/comparison')
@login_required
def comparison():
    from .services.comparison import compare
    assessments = db.session.scalars(assessments_query().order_by(Assessment.id.desc())).all()
    baseline = retest = result = None
    if request.args.get('baseline_id') or request.args.get('retest_id'):
        try:
            baseline = get_assessment(int(request.args.get('baseline_id', '')))
            retest = get_assessment(int(request.args.get('retest_id', '')))
        except ValueError:
            abort(400)
        if baseline.id == retest.id or baseline.owner_id != retest.owner_id or retest.baseline_id != baseline.id:
            flash('Select a baseline and its linked retest under the same owner.', 'warning')
        else:
            result = compare(baseline, retest)
    return render_template('comparison.html', assessments=assessments, baseline=baseline, retest=retest, result=result)


@web.get('/reports')
@login_required
def reports():
    a = selected_assessment()
    query = db.select(Report).join(Assessment)
    if not current_user.is_admin:
        query = query.where(Assessment.owner_id == current_user.id)
    if a:
        query = query.where(Assessment.id == a.id)
    # Defer binary content while listing report history.
    from sqlalchemy.orm import defer
    pagination = db.paginate(query.options(defer(Report.content)).order_by(Report.id.desc()), per_page=20, max_per_page=20)
    return render_template('reports.html', assessment=a, assessments=db.session.scalars(assessments_query()).all(), pagination=pagination, report_form=ReportForm())


@web.post('/assessments/<int:identifier>/reports')
@login_required
def report_generate(identifier):
    a = get_assessment(identifier)
    form = ReportForm()
    if form.validate_on_submit():
        from .services.reports import generate, snapshot
        try:
            data = snapshot(a, current_user)
            content, digest = generate(data, form.format.data)
        except ValueError as exc:
            flash(str(exc), 'danger')
        else:
            report = Report(assessment_id=a.id, generated_by_id=current_user.id, format=form.format.data,
                            filename=f'vulntrack-assessment-{a.id}-{now():%Y%m%d-%H%M%S}.{form.format.data}',
                            snapshot=data, content=content, sha256=digest)
            db.session.add(report)
            db.session.flush()
            audit('report_generated', a, f'Report #{report.id} ({report.format}); SHA-256 {digest}')
            db.session.commit()
            flash('Report saved to history. Download your snapshot below.', 'success')
            return redirect(url_for('web.reports', assessment_id=a.id))
    else:
        form_errors(form)
    return redirect(url_for('web.assessment_detail', identifier=a.id))


@web.get('/reports/<int:identifier>/download')
@login_required
def report_download(identifier):
    report = db.session.get(Report, identifier) or abort(404)
    get_assessment(report.assessment_id)
    audit('report_downloaded', report.assessment, f'Report #{report.id}')
    db.session.commit()
    mimetype = 'application/pdf' if report.format == 'pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return send_file(BytesIO(report.content), mimetype=mimetype, as_attachment=True, download_name=report.filename, max_age=0)


@web.get('/scans')
@login_required
def scans():
    query = db.select(ScanJob).join(Assessment)
    if not current_user.is_admin:
        query = query.where(Assessment.owner_id == current_user.id)
    pagination = db.paginate(query.order_by(ScanJob.id.desc()), per_page=20, max_per_page=20)
    return render_template('scans.html', pagination=pagination)


@web.post('/assessments/<int:identifier>/scan')
@login_required
def scan_start(identifier):
    a = get_assessment(identifier)
    form = ScanForm()
    if form.validate_on_submit():
        from .services.authorization import active_approval
        if form.active.data and (not current_app.config.get('ZAP_ENABLED') or not active_approval(a)):
            flash('Active tests require separate, current administrator approval and the ZAP engine.', 'warning')
            return redirect(url_for('web.automatic_scan', identifier=a.id))
        if not a.scan_approved:
            flash('An administrator must review authorization and approve this target first.', 'warning')
        elif db.session.scalar(db.select(ScanJob).where(ScanJob.assessment_id == a.id, ScanJob.status.in_(['queued', 'running']))):
            flash('This assessment already has an active scan.', 'warning')
        else:
            from .services.zap_engine import available
            if current_app.config.get('ZAP_ENABLED') and not available(current_app):
                flash('Automatic scanning needs the free ZAP engine. Ask the administrator to run install-zap.ps1 and restart the worker.', 'warning')
                return redirect(url_for('web.automatic_scan', identifier=a.id))
            active_count = db.session.scalar(db.select(db.func.count()).select_from(ScanJob).join(Assessment).where(Assessment.owner_id == a.owner_id, ScanJob.status.in_(['queued', 'running'])))
            if active_count >= 2:
                flash('Two scans are already active for this owner. Wait or cancel one.', 'warning')
            else:
                job = ScanJob(assessment_id=a.id, requested_by_id=current_user.id, target_url=a.target_url,
                              result={'profile': 'active' if form.active.data else 'passive', 'stage': 'Queued for the background worker'})
                db.session.add(job)
                audit('scan_queued', a, 'Automatic scan requested: ' + ('approved active SQLi/XSS testing' if form.active.data else 'low-impact passive analysis'))
                try:
                    db.session.commit()
                except IntegrityError:
                    db.session.rollback()
                    flash('Another request already queued a scan for this assessment.', 'warning')
                    return redirect(url_for('web.scans'))
                return redirect(url_for('web.scan_progress', identifier=job.id))
    else:
        form_errors(form)
    return redirect(url_for('web.automatic_scan', identifier=a.id))


@web.get('/scans/<int:identifier>')
@login_required
def scan_progress(identifier):
    job = db.session.get(ScanJob, identifier) or abort(404)
    assessment = get_assessment(job.assessment_id)
    return render_template('scan_progress.html', assessment=assessment, job=job)


@web.get('/scans/<int:identifier>/status')
@login_required
def scan_status(identifier):
    job = db.session.get(ScanJob, identifier) or abort(404)
    get_assessment(job.assessment_id)
    from flask import jsonify
    return jsonify(status=job.status, findings_count=job.findings_count, pages_checked=job.pages_checked,
                   error=job.error, stage=(job.result or {}).get('stage', ''), cancel_requested=job.cancel_requested)


@web.post('/scans/<int:identifier>/cancel')
@login_required
def scan_cancel(identifier):
    job = db.session.get(ScanJob, identifier) or abort(404)
    a = get_assessment(job.assessment_id)
    if job.status in ('queued', 'running'):
        job.cancel_requested = True
        if job.status == 'queued':
            job.status, job.finished_at = 'cancelled', now()
        audit('scan_cancel_requested', a, f'Scan #{job.id}')
        db.session.commit()
        flash('Cancellation requested. A running request may take a few seconds to finish.', 'info')
    return redirect(url_for('web.scan_progress', identifier=job.id))


@web.route('/account', methods=['GET', 'POST'])
@login_required
def account():
    form = PasswordForm()
    if form.validate_on_submit():
        if not current_user.check_password(form.current_password.data):
            form.current_password.errors.append('Current password is incorrect.')
        else:
            current_user.set_password(form.password.data)
            audit('password_changed')
            db.session.commit()
            logout_user()
            session.clear()
            flash('Password changed. Sign in with your new password.', 'success')
            return redirect(url_for('web.login'))
    return render_template('account.html', form=form)


@web.route('/users', methods=['GET', 'POST'])
@admin_required
def users():
    form = UserForm()
    if form.validate_on_submit():
        email = form.email.data.strip().lower()
        if db.session.scalar(db.select(User).where(User.email == email)):
            form.email.errors.append('An account already uses that email address.')
        else:
            user = User(name=form.name.data.strip(), email=email, role=form.role.data)
            user.set_password(form.password.data)
            db.session.add(user)
            try:
                db.session.flush()
                audit('user_created', details=f'Account #{user.id}; role {user.role}')
                db.session.commit()
            except IntegrityError:
                db.session.rollback()
                form.email.errors.append('An account already uses that email address.')
            else:
                flash('Account created. Share its initial password through your approved channel.', 'success')
                return redirect(url_for('web.users'))
    return render_template('users.html', form=form, users=db.session.scalars(db.select(User).order_by(User.created_at)).all())


@web.post('/users/<int:identifier>/toggle')
@admin_required
def user_toggle(identifier):
    user = db.session.get(User, identifier) or abort(404)
    if user.id == current_user.id:
        flash('You cannot disable your own account.', 'warning')
    else:
        user.enabled = not user.enabled
        user.auth_token = secrets.token_hex(24)
        audit('user_enabled' if user.enabled else 'user_disabled', details=f'Account #{user.id}')
        db.session.commit()
        flash('Account access updated and previous sessions invalidated.', 'success')
    return redirect(url_for('web.users'))


@web.get('/audit')
@admin_required
def audit_log():
    pagination = db.paginate(db.select(AuditLog).order_by(AuditLog.id.desc()), per_page=30, max_per_page=30)
    return render_template('audit.html', pagination=pagination)


@web.route('/findings/<int:identifier>/delete', methods=['GET', 'POST'])
@login_required
def finding_delete(identifier):
    finding = get_finding(identifier)
    form = DeleteForm()
    if form.validate_on_submit():
        if form.confirmation.data != finding.title or not current_user.check_password(form.current_password.data):
            flash('The title or current password did not match. Nothing was deleted.', 'danger')
        else:
            assessment = finding.assessment
            db.session.execute(db.update(Finding).where(Finding.duplicate_of_id == finding.id).values(duplicate_of_id=None))
            audit('finding_deleted', assessment, f'Finding #{finding.id}: {finding.title}')
            db.session.delete(finding)
            db.session.commit()
            flash('Finding deleted. Existing report snapshots retain their historical content.', 'success')
            return redirect(url_for('web.findings', assessment_id=assessment.id))
    return render_template('delete.html', form=form, record=finding, kind='finding',
                           back_url=url_for('web.finding_detail', identifier=finding.id))


@web.route('/assessments/<int:identifier>/delete', methods=['GET', 'POST'])
@login_required
def assessment_delete(identifier):
    assessment = get_assessment(identifier)
    form = DeleteForm()
    if form.validate_on_submit():
        active = db.session.scalar(db.select(ScanJob).where(ScanJob.assessment_id == assessment.id, ScanJob.status.in_(['queued', 'running'])))
        linked = db.session.scalar(db.select(Assessment).where(Assessment.baseline_id == assessment.id))
        if form.confirmation.data != assessment.title or not current_user.check_password(form.current_password.data):
            flash('The title or current password did not match. Nothing was deleted.', 'danger')
        elif active or linked:
            flash('Finish or cancel active scans and remove linked retests before deleting this assessment.', 'warning')
        else:
            original_id, title = assessment.id, assessment.title
            db.session.execute(db.update(Finding).where(Finding.assessment_id == original_id).values(duplicate_of_id=None))
            db.session.execute(db.delete(Report).where(Report.assessment_id == original_id))
            db.session.execute(db.delete(ImportBatch).where(ImportBatch.assessment_id == original_id))
            db.session.execute(db.delete(ActiveScanApproval).where(ActiveScanApproval.assessment_id == original_id))
            db.session.execute(db.delete(ScanJob).where(ScanJob.assessment_id == original_id))
            db.session.execute(db.delete(Finding).where(Finding.assessment_id == original_id))
            db.session.execute(db.update(AuditLog).where(AuditLog.assessment_id == original_id).values(assessment_id=None))
            audit('assessment_deleted', details=f'Assessment #{original_id}: {title}; related reports and findings removed')
            db.session.execute(db.delete(Assessment).where(Assessment.id == original_id))
            db.session.commit()
            flash('Assessment and its findings, scan history, and report snapshots were deleted.', 'success')
            return redirect(url_for('web.assessments'))
    return render_template('delete.html', form=form, record=assessment, kind='assessment',
                           back_url=url_for('web.assessment_detail', identifier=assessment.id))
