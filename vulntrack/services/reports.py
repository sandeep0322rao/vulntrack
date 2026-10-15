import hashlib
from collections import Counter
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab import __file__ as reportlab_path
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

from ..constants import CATEGORIES, OPEN_STATUSES, STATUSES
from ..models import now
from .risk import priority

FORMULA = 'Risk = severity × 0.40 + exploitability × 0.25 + business impact × 0.25 + exposure × 0.10. Inputs use a 0–100 scale. This project score is not CVSS.'
LIMITATIONS = 'Automated alerts require assessor review. This report covers the recorded assessment scope and available evidence; it does not certify that the target is free of vulnerabilities. Synthetic records are demonstrations only.'


def snapshot(assessment, user):
    findings = sorted(assessment.findings, key=lambda f: (-f.risk_score, f.id))
    if len(findings) > 500:
        raise ValueError('A report is limited to 500 findings. Split larger assessments before reporting.')
    keys = ['id', 'title', 'url', 'method', 'parameter', 'cwe', 'severity', 'confidence', 'category', 'description', 'evidence', 'remediation', 'source', 'exploitability', 'business_impact', 'exposure', 'risk_score', 'status', 'assigned_to', 'resolution_note', 'fingerprint', 'duplicate_of_id']
    return {
        'schema_version': 1, 'generated_at_utc': now().isoformat(timespec='seconds'), 'generated_by': user.name,
        'assessment': {k: getattr(assessment, k) for k in ['id', 'title', 'client', 'target_url', 'scope', 'authorization_ref', 'kind', 'baseline_id']},
        'formula': FORMULA, 'limitations': LIMITATIONS,
        'summary': {'total': len(findings), 'open': sum(f.status in OPEN_STATUSES for f in findings), 'severity': dict(Counter(f.severity for f in findings))},
        'findings': [{**{k: getattr(f, k) for k in keys}, 'retested_by': f.retested_by.name if f.retested_by else '',
                      'retested_at_utc': f.retested_at.isoformat() if f.retested_at else ''} for f in findings],
    }


def excel_text(value):
    # Assign as explicit string cells, even if input begins with =, +, -, or @.
    return ILLEGAL_CHARACTERS_RE.sub('', str(value or ''))[:32767]


def xlsx_bytes(data):
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    summary.append(['VulnTrack assessment report', 'Value'])
    for k, v in data['assessment'].items():
        summary.append([k.replace('_', ' ').title(), '' if v is None else str(v)])
    for k in ['generated_at_utc', 'generated_by', 'formula', 'limitations']:
        summary.append([k.replace('_', ' ').title(), data[k]])
    summary.append(['Total findings', data['summary']['total']])
    summary.append(['Open risk findings (including accepted risk)', data['summary']['open']])
    findings = workbook.create_sheet('Findings')
    keys = ['id', 'title', 'severity', 'confidence', 'risk_score', 'category', 'url', 'method', 'parameter', 'cwe', 'status', 'source', 'description', 'evidence', 'remediation', 'exploitability', 'business_impact', 'exposure', 'duplicate_of_id', 'fingerprint']
    findings.append([k.replace('_', ' ').title() for k in keys])
    remediation = workbook.create_sheet('Remediation')
    rkeys = ['id', 'title', 'status', 'assigned_to', 'resolution_note', 'retested_by', 'retested_at_utc']
    remediation.append([k.replace('_', ' ').title() for k in rkeys])
    for f in data['findings']:
        findings.append([f[k] for k in keys])
        remediation.append([f[k] for k in rkeys])
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = excel_text(cell.value)
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row == 1:
                    cell.fill = PatternFill('solid', fgColor='173D39')
                    cell.font = Font(color='FFFFFF', bold=True)
                elif cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F1F6F4')
        for col in sheet.columns:
            label = str(col[0].value or '')
            sheet.column_dimensions[col[0].column_letter].width = 50 if label in ('Title', 'Description', 'Evidence', 'Remediation', 'Value', 'Resolution Note', 'Url') else 24
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.page_states = []

    def showPage(self):
        self.page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        pages = len(self.page_states)
        for state in self.page_states:
            self.__dict__.update(state)
            self.setFont('VT', 8)
            self.setFillColor(colors.HexColor('#65756F'))
            self.drawString(42, 27, 'VulnTrack | Confidential assessment report')
            self.drawRightString(self._pagesize[0] - 42, 27, f'Page {self._pageNumber} of {pages}')
            super().showPage()
        super().save()


def pdf_bytes(data):
    fonts = Path(reportlab_path).parent / 'fonts'
    if 'VT' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('VT', str(fonts / 'Vera.ttf')))
        pdfmetrics.registerFont(TTFont('VT-Bold', str(fonts / 'VeraBd.ttf')))
        pdfmetrics.registerFontFamily('VT', normal='VT', bold='VT-Bold', italic='VT', boldItalic='VT-Bold')
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = 'VT'
    styles['BodyText'].fontSize = 9
    styles['BodyText'].leading = 14
    styles['BodyText'].wordWrap = 'CJK'
    styles['Heading1'].fontName = 'VT-Bold'
    styles['Heading1'].spaceBefore = 16
    styles['Heading2'].fontName = 'VT-Bold'
    styles['Heading2'].fontSize = 13
    stream = BytesIO()
    document = SimpleDocTemplate(stream, pagesize=(595.28, 841.89), leftMargin=42, rightMargin=42, topMargin=45, bottomMargin=49,
                                 title=f"VulnTrack assessment {data['assessment']['id']}", author='VulnTrack')
    story = []

    def p(text, style='BodyText'):
        return Paragraph(escape(str(text)).replace('\n', '<br/>'), styles[style])

    def section(title, value):
        story.append(p(title, 'Heading2'))
        # Bound each paragraph to avoid unbreakable evidence blocks.
        value = str(value or 'Not recorded')
        for i in range(0, len(value), 1800):
            story.append(p(value[i:i + 1800]))

    a = data['assessment']
    story.extend([Spacer(1, 55), p('VulnTrack', 'Title'), p('Vulnerability Assessment Report', 'Heading1'), Spacer(1, 20),
                  p(a['title'], 'Heading1'), p(a['client']), p(a['target_url']), Spacer(1, 24),
                  p(f"Prepared by {data['generated_by']}"), p(f"Generated {data['generated_at_utc']} UTC"),
                  Spacer(1, 30), p(data['limitations']), PageBreak(), p('Executive summary', 'Heading1')])
    story.append(p(f"{data['summary']['total']} recorded findings. {data['summary']['open']} findings contribute to open risk, including accepted risks. Potential duplicates remain visible and are counted until reviewed."))
    rows = [[p('Severity'), p('Count')]] + [[p(k), p(v)] for k, v in data['summary']['severity'].items()]
    table = Table(rows, colWidths=[360, 140], hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EAF1EF')), ('GRID', (0, 0), (-1, -1), .5, colors.HexColor('#D9D9D9')), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('TOPPADDING', (0, 0), (-1, -1), 9), ('BOTTOMPADDING', (0, 0), (-1, -1), 9)]))
    story.append(table)
    section('Assessment scope', a['scope'])
    section('Authorization reference', a['authorization_ref'])
    section('Methodology and coverage', 'Findings come from the source recorded on each item: manual review, Excel/CSV uploads, imported ZAP exports, or automatic scans. The default scan analyzes one unauthenticated response with free ZAP passive rules and built-in checks. Separately approved active scans enable only low-strength SQL injection and reflected XSS rules on the selected endpoint, without crawling. No mode provides a complete authenticated or business logic assessment. OWASP WSTG is a testing methodology reference: https://owasp.org/www-project-web-security-testing-guide/.')
    section('Risk prioritization', data['formula'])
    section('Priority bands', 'Critical 80–100; High 60–<80; Medium 40–<60; Low 0–<40. Severity presets: Critical 100, High 80, Medium 55, Low 30, Informational 10.')
    if not data['findings']:
        section('Findings', 'No findings are recorded. This is not evidence that the target is secure.')
    for f in data['findings']:
        story.append(PageBreak())
        story.append(p(f"Finding {f['id']} · {f['title']}", 'Heading1'))
        story.append(p(f"{f['severity']} severity | {f['risk_score']:.2f}/100 | {priority(f['risk_score'])} priority | {STATUSES[f['status']]}"))
        section('Affected resource', f"{f['method']} {f['url']}\nParameter: {f['parameter'] or 'None recorded'}")
        section('Classification and source', f"{f['category']} · {CATEGORIES.get(f['category'], '')}\nCWE: {f['cwe'] or 'Not recorded'}\nSource: {f['source']} | Confidence: {f['confidence']}")
        section('Description and impact', f['description'])
        section('Evidence', f['evidence'])
        section('Remediation guidance', f['remediation'])
        section('Risk inputs', f"Exploitability {f['exploitability']}; business impact {f['business_impact']}; exposure {f['exposure']}. {data['formula']}")
        section('Remediation and retest', f"Assigned to: {f['assigned_to'] or 'Unassigned'}\n{f['resolution_note']}\nRetested by: {f['retested_by'] or 'Not retested'}\nRetest date (UTC): {f['retested_at_utc'] or 'Not recorded'}")
        if f['duplicate_of_id']:
            section('Duplicate review', f"Potential duplicate of finding {f['duplicate_of_id']}. Retained for assessor review.")
    document.build(story, canvasmaker=NumberedCanvas)
    return stream.getvalue()


def generate(data, format):
    content = pdf_bytes(data) if format == 'pdf' else xlsx_bytes(data)
    return content, hashlib.sha256(content).hexdigest()
