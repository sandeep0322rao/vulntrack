"""Strict, bounded spreadsheet intake. No formulas or macros are evaluated."""
import csv
import math
import re
from io import BytesIO, StringIO
from zipfile import BadZipFile, ZipFile
from itertools import islice
from defusedxml.common import DefusedXmlException
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..constants import CATEGORIES
from ..forms import valid_http_url
from .intake import MAX_FINDINGS, clean

HEADERS = ['Title', 'Severity', 'URL', 'Description', 'Evidence', 'Remediation', 'Method', 'Parameter', 'CWE', 'Category', 'Exploitability', 'Business Impact', 'Exposure']
ALIASES = {
    'title': ['title', 'finding', 'finding title', 'issue', 'issue name', 'vulnerability', 'vulnerability name', 'name'],
    'severity': ['severity', 'risk level', 'risk rating', 'risk'],
    'url': ['url', 'affected url', 'target url', 'affected resource', 'uri'],
    'description': ['description', 'description and impact', 'impact', 'details'],
    'evidence': ['evidence', 'proof', 'proof of concept', 'poc'],
    'remediation': ['remediation', 'remediation guidance', 'recommendation', 'recommendations', 'solution', 'fix'],
    'method': ['method', 'http method'], 'parameter': ['parameter', 'param'],
    'cwe': ['cwe', 'cwe id', 'cwe identifier'], 'category': ['category', 'owasp', 'owasp category', 'owasp top 10'],
    'exploitability': ['exploitability'], 'business_impact': ['business impact', 'business_impact'], 'exposure': ['exposure'],
}


def header_key(value):
    return re.sub(r'[^a-z0-9]+', ' ', str(value or '').lower()).strip()


def normalize_record(row, target_url, number, source):
    title = str(row.get('title') or '').strip()
    if not title or len(title) > 240:
        raise ValueError('Title is required and must be at most 240 characters.')
    severity = str(row.get('severity') or '').strip().lower()
    severities = {'critical': 'Critical', 'high': 'High', 'medium': 'Medium', 'moderate': 'Medium', 'low': 'Low', 'info': 'Informational', 'informational': 'Informational'}
    if severity not in severities:
        raise ValueError('Severity must be Critical, High, Medium, Low, or Informational.')
    url = str(row.get('url') or target_url).strip()
    valid_http_url(url)
    method = str(row.get('method') or 'GET').strip().upper()
    if method not in ('GET', 'POST', 'PUT', 'PATCH', 'DELETE', 'HEAD', 'OPTIONS'):
        raise ValueError('HTTP method is not supported.')
    record = {'title': clean(title, 240), 'severity': severities[severity], 'url': url, 'method': method,
              'description': clean(row.get('description') or 'Imported observation. Review and document its impact.'),
              'evidence': clean(row.get('evidence')), 'remediation': clean(row.get('remediation') or 'Review the issue and record appropriate remediation guidance.'),
              'parameter': clean(row.get('parameter'), 160), 'cwe': clean(row.get('cwe'), 40),
              'source': source, 'confidence': 'Needs review', '_row': number}
    if not record['title']:
        raise ValueError('Title must contain readable text.')
    for key, default in [('exploitability', 50), ('business_impact', 50), ('exposure', 75)]:
        value = row.get(key)
        if value is None or str(value).strip() == '':
            record[key] = default
            continue
        try:
            numeric = float(value)
            if isinstance(value, bool) or not math.isfinite(numeric) or not numeric.is_integer() or not 0 <= numeric <= 100:
                raise ValueError()
        except (ValueError, TypeError, OverflowError):
            raise ValueError(f'{key.replace("_", " ").title()} must be a whole number from 0 to 100.') from None
        record[key] = int(numeric)
    category = str(row.get('category') or '').strip()
    if category:
        if category in CATEGORIES:
            record['category'] = category
        else:
            match = re.match(r'^(A\d{2}:2025)(?:\s|[·-]|$)', category, re.I)
            if not match or match[1].upper() not in CATEGORIES:
                raise ValueError('Category must use an OWASP 2025 identifier, such as A05:2025, or be blank.')
            record['category'] = match[1].upper()
        record['mapping_reason'] = 'Classification supplied in uploaded spreadsheet; assessor review required.'
    return record


def parse_rows(rows, target_url, source):
    aliases = {header_key(alias): key for key, values in ALIASES.items() for alias in values}
    records, errors, warnings = [], [], []
    try:
        header = next(rows)
    except StopIteration:
        raise ValueError('The selected worksheet is empty.') from None
    if len(header) > 80:
        raise ValueError('A spreadsheet may contain at most 80 columns.')
    columns = {}
    for index, value in enumerate(header):
        key = aliases.get(header_key(value))
        if key:
            if key in columns.values():
                raise ValueError(f'More than one column maps to {key}. Remove duplicate columns before uploading.')
            columns[index] = key
    missing = {'title', 'severity'} - set(columns.values())
    if missing:
        raise ValueError('Required columns not found: ' + ', '.join(sorted(missing)) + '. Download the template for the supported format.')
    if 'url' not in columns.values():
        warnings.append('No URL column was found. The assessment target is used for each row.')
    if 'remediation' not in columns.values() or 'description' not in columns.values():
        warnings.append('Some details are missing. Imported findings will require assessor review.')
    for number, values in enumerate(rows, start=2):
        if number > MAX_FINDINGS + 1:
            raise ValueError(f'Import is limited to {MAX_FINDINGS} data rows. Split this file before uploading.')
        if not any(v is not None and str(v).strip() for v in values):
            continue
        if len(values) > 80:
            errors.append({'row': number, 'message': 'Too many columns.'})
            continue
        try:
            if any(isinstance(v, Formula) for v in values):
                raise ValueError('Formula cells are not imported. Replace formulas with plain values.')
            row = {key: values[index] if index < len(values) else None for index, key in columns.items()}
            records.append(normalize_record(row, target_url, number, source))
        except ValueError as exc:
            errors.append({'row': number, 'message': str(exc)})
    if not records and not errors:
        raise ValueError('No findings were found below the header row.')
    warnings.append('Imported findings start Open with confidence Needs review. Risk scores are recalculated; uploaded status and score columns are not trusted.')
    return records, errors, warnings


class Formula(str):
    pass


def spreadsheet_records(payload, extension, target_url):
    if extension == 'csv':
        try:
            text = payload.decode('utf-8-sig')
            csv.field_size_limit(36000)
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            return parse_rows(iter(csv.reader(StringIO(text), dialect)), target_url, 'CSV import')
        except (UnicodeError, csv.Error):
            raise ValueError('CSV must use UTF-8 encoding and valid quoted columns.') from None
    if extension != 'xlsx':
        raise ValueError('Use an .xlsx workbook or UTF-8 .csv file. Convert older .xls files to .xlsx.')
    try:
        with ZipFile(BytesIO(payload)) as archive:
            info = archive.infolist()
            if len(info) > 2000 or sum(f.file_size for f in info) > 32 * 1024 * 1024:
                raise ValueError('The expanded workbook exceeds the 32 MB safety limit.')
            if any(f.filename.endswith('vbaProject.bin') for f in info):
                raise ValueError('Macro-enabled content is not accepted.')
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook['Findings'] if 'Findings' in workbook.sheetnames else workbook.worksheets[0]
            if sheet.max_column and sheet.max_column > 80:
                raise ValueError('A worksheet may contain at most 80 columns.')
            # Do not trust the worksheet's claimed row count, and never iterate an unbounded range.
            sheet.reset_dimensions()
            rows = ([Formula(cell.value) if cell.data_type == 'f' else cell.value for cell in row]
                    for row in islice(sheet.iter_rows(max_col=80), MAX_FINDINGS + 2))
            result = parse_rows(iter(rows), target_url, 'Excel import')
            result[2].insert(0, f'Worksheet selected: {sheet.title}.')
            return result
        finally:
            workbook.close()
    except (BadZipFile, InvalidFileException, KeyError, IndexError, TypeError, OSError, DefusedXmlException, ParseError):
        raise ValueError('The file is not a readable .xlsx workbook. Save it as a normal Excel workbook and try again.') from None
